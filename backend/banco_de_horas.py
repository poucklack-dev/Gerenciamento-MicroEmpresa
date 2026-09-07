"""Banco de horas mensal: configuração explícita, consolidação e auditoria."""
import io, re
from calendar import monthrange
from datetime import date, datetime, timedelta
from flask import Blueprint, jsonify, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename
from backend.core.database import get_conn

banco_horas_bp = Blueprint("banco_horas", __name__, url_prefix="/api/banco_horas")

def parse_minutes(value, signed=False):
    if isinstance(value, int): return value
    if isinstance(value, float): return round(value * 60)
    text = str(value or "").strip().lower().replace(" ", "")
    match = re.fullmatch(r"([+-]?)(\d+)(?:h|:)?(\d{1,2})?", text)
    if not match: raise ValueError("Use horas como 8, 8:30, 8h30 ou 172:00.")
    sign = -1 if match.group(1) == "-" else 1
    hours, minutes = int(match.group(2)), int(match.group(3) or 0)
    if minutes > 59: raise ValueError("Minutos devem estar entre 00 e 59.")
    result = sign * (hours * 60 + minutes)
    if not signed and result < 0: raise ValueError("A duração não pode ser negativa.")
    return result

def format_minutes(value, show_sign=False):
    value = int(value or 0); sign = "-" if value < 0 else ("+" if show_sign and value > 0 else "")
    value = abs(value)
    return f"{sign}{value // 60}h{value % 60:02d}"

def calculate_expected(daily_minutes=None, days=None, adjustments=None, direct_minutes=None):
    if direct_minutes is not None:
        result = int(direct_minutes)
    else:
        result = int(daily_minutes or 0) * int(days or 0) + sum(int(x) for x in (adjustments or []))
    if result < 0:
        raise ValueError("A carga prevista final não pode ser negativa.")
    return result

def official_worked(source, records_minutes, manual_minutes=None):
    """Seleciona uma única fonte oficial; manual e registros nunca são somados."""
    return int(manual_minutes or 0) if source == "manual" else int(records_minutes or 0)

def calculate_balance(expected_minutes, worked_minutes, excluded=False):
    balance = int(worked_minutes or 0) - int(expected_minutes or 0)
    return {"saldo_minutos": balance, "extras_minutos": 0 if excluded else max(balance, 0),
            "negativas_minutos": 0 if excluded else max(-balance, 0), "considerar_nos_totais": not excluded}

def calculate_record_minutes(start_time, end_time, interval_minutes=0):
    if not start_time or not end_time:
        return 0
    elapsed = round((end_time - start_time).total_seconds() / 60)
    if elapsed < 0:
        elapsed += 24 * 60
    return max(elapsed - int(interval_minutes or 0), 0)

def competence(value):
    try: parsed = datetime.strptime(value, "%Y-%m").date()
    except (TypeError, ValueError): raise ValueError("Período inválido. Use AAAA-MM.")
    return parsed.replace(day=1)

def period_end(start): return start.replace(day=monthrange(start.year, start.month)[1])
def user_id(): return (session.get("usuarios") or {}).get("id")

def ensure_period(cur, start):
    cur.execute("INSERT INTO banco_horas_periodos(competencia,criado_por) VALUES(%s,%s) ON CONFLICT(competencia) DO UPDATE SET competencia=EXCLUDED.competencia RETURNING id", (start,user_id()))
    return cur.fetchone()[0]

def audit(cur, period_id, collaborator_id, action, old=None, new=None, reason=None):
    cur.execute("INSERT INTO banco_horas_auditoria(periodo_id,colaborador_id,usuario_id,acao,valor_anterior,valor_novo,motivo) VALUES(%s,%s,%s,%s,%s,%s,%s)", (period_id,collaborator_id,user_id(),action,None if old is None else str(old),None if new is None else str(new),reason))

SUMMARY_SQL = """
WITH registros AS (
 SELECT colaborador_id,
   COALESCE(SUM(CASE WHEN situacao_dia='presente' AND hora_entrada IS NOT NULL AND hora_saida IS NOT NULL
     THEN GREATEST((EXTRACT(EPOCH FROM (hora_saida-hora_entrada))/60)::integer-COALESCE(intervalo_minutos,0),0) ELSE 0 END),0)::integer trabalhados,
   COUNT(DISTINCT data_registro) FILTER (WHERE situacao_dia='presente')::integer presentes,
   COUNT(DISTINCT data_registro) FILTER (WHERE situacao_dia IN ('ausente','ausencia_justificada'))::integer ausentes
 FROM ponto WHERE data_registro BETWEEN %s AND %s GROUP BY colaborador_id
)
SELECT c.id,c.nome,c.cargo,COALESCE(c.setor,'Sem setor'),COALESCE(ch.nome,''),
 COALESCE(bc.horas_previstas_minutos,bp.carga_prevista_padrao_minutos,0)::integer previstas,
 CASE WHEN COALESCE(bc.origem_horas,'registros')='manual' THEN COALESCE(bc.horas_trabalhadas_manual_minutos,0) ELSE COALESCE(r.trabalhados,0) END::integer trabalhadas,
 COALESCE(bc.origem_horas,'registros'),
 CASE WHEN COALESCE(bc.origem_frequencia,'registros')='manual' THEN COALESCE(bc.dias_presentes_manual,0) ELSE COALESCE(r.presentes,0) END::integer presentes,
 CASE WHEN COALESCE(bc.origem_frequencia,'registros')='manual' THEN COALESCE(bc.dias_ausentes_manual,0) ELSE COALESCE(r.ausentes,0) END::integer ausentes,
 COALESCE(bc.excluido_periodo,FALSE),bc.motivo_exclusao,COALESCE(bc.origem_frequencia,'registros')
FROM colaboradores c
LEFT JOIN colaboradores ch ON ch.id=c.chefe_direto_id
LEFT JOIN banco_horas_periodos bp ON bp.competencia=%s
LEFT JOIN banco_horas_colaborador bc ON bc.periodo_id=bp.id AND bc.colaborador_id=c.id
LEFT JOIN registros r ON r.colaborador_id=c.id
WHERE LOWER(COALESCE(c.status,'ativo'))='ativo'
"""

def get_summary(cur, start, args):
    sql=SUMMARY_SQL; values=[start,period_end(start),start]
    if args.get("colaborador_id"): sql+=" AND c.id=%s"; values.append(int(args["colaborador_id"]))
    if args.get("setor"): sql+=" AND COALESCE(c.setor,'Sem setor')=%s"; values.append(args["setor"])
    sql+=" ORDER BY c.nome"; cur.execute(sql,values)
    result=[]
    for row in cur.fetchall():
        excluded=bool(row[10]); balance=calculate_balance(row[5],row[6],excluded); saldo=balance["saldo_minutos"]
        result.append({"id":row[0],"nome":row[1],"cargo":row[2] or "—","setor":row[3],"chefe_direto":row[4],"previstas_minutos":row[5],"trabalhadas_minutos":row[6],**balance,"origem_horas":row[7],"dias_presentes":row[8],"dias_ausentes":row[9],"excluido_periodo":excluded,"motivo_exclusao":row[11],"origem_frequencia":row[12],"previstas":format_minutes(row[5]),"trabalhadas":format_minutes(row[6]),"saldo":format_minutes(saldo,True),"situacao":"Anulado" if excluded else ("Positivo" if saldo>0 else "Negativo" if saldo<0 else "Regular")})
    balance=args.get("saldo"); status=args.get("status","validos")
    if balance in {"positivo","negativo","regular"}: result=[x for x in result if x["situacao"].lower()==balance]
    if status=="validos": result=[x for x in result if not x["excluido_periodo"]]
    elif status=="anulados": result=[x for x in result if x["excluido_periodo"]]
    return result

@banco_horas_bp.get("/resumo")
def summary():
    try: start=competence(request.args.get("competencia"))
    except ValueError as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        rows=get_summary(cur,start,request.args); valid=[r for r in rows if not r["excluido_periodo"]]
        cur.execute("SELECT metodo_carga,carga_diaria_minutos,dias_considerados,carga_base_minutos,carga_prevista_padrao_minutos,observacao FROM banco_horas_periodos WHERE competencia=%s",(start,)); p=cur.fetchone()
        totals={k:sum(r[k] for r in valid) for k in ("previstas_minutos","trabalhadas_minutos","saldo_minutos","extras_minutos","negativas_minutos","dias_presentes","dias_ausentes")}
        totals.update({"previstas":format_minutes(totals["previstas_minutos"]),"trabalhadas":format_minutes(totals["trabalhadas_minutos"]),"saldo":format_minutes(totals["saldo_minutos"],True)})
        return jsonify({"competencia":start.strftime("%Y-%m"),"colaboradores":rows,"totais":totals,"configuracao":None if not p else {"metodo":p[0],"carga_diaria_minutos":p[1],"dias_considerados":p[2],"carga_base_minutos":p[3],"prevista_minutos":p[4],"prevista":format_minutes(p[4]),"observacao":p[5]}})
    finally: cur.close(); conn.close()

@banco_horas_bp.get("/colaboradores")
def collaborators():
    conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("SELECT id,nome,cargo,COALESCE(setor,'Sem setor') FROM colaboradores WHERE LOWER(COALESCE(status,'ativo'))='ativo' ORDER BY nome")
        return jsonify([{"id":r[0],"nome":r[1],"cargo":r[2],"setor":r[3]} for r in cur.fetchall()])
    finally: cur.close(); conn.close()

@banco_horas_bp.post("/configurar-carga")
def configure_load():
    data=request.get_json(silent=True) or {}
    try:
        start=competence(data.get("competencia")); method=data.get("metodo","direto")
        adjustments=[{"minutos":parse_minutes(a.get("horas"),True),"motivo":str(a.get("motivo","")).strip()} for a in data.get("ajustes",[])]
        if method=="calculado":
            daily=parse_minutes(data.get("carga_diaria")); days=int(data.get("dias_considerados",0)); base=daily*days; final=calculate_expected(daily,days,[a["minutos"] for a in adjustments])
        elif method=="direto": daily=None; days=None; base=parse_minutes(data.get("total_previsto")); final=calculate_expected(direct_minutes=base)
        else: raise ValueError("Método de carga inválido.")
        if final<0: raise ValueError("A carga prevista final não pode ser negativa.")
    except (ValueError,TypeError) as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        pid=ensure_period(cur,start); cur.execute("SELECT carga_prevista_padrao_minutos FROM banco_horas_periodos WHERE id=%s",(pid,)); old=cur.fetchone()[0]
        cur.execute("UPDATE banco_horas_periodos SET metodo_carga=%s,carga_diaria_minutos=%s,dias_considerados=%s,carga_base_minutos=%s,carga_prevista_padrao_minutos=%s,observacao=%s,atualizado_em=NOW() WHERE id=%s",(method,daily,days,base,final,data.get("observacao"),pid))
        cur.execute("DELETE FROM banco_horas_ajustes WHERE periodo_id=%s",(pid,))
        for a in adjustments: cur.execute("INSERT INTO banco_horas_ajustes(periodo_id,minutos,motivo) VALUES(%s,%s,%s)",(pid,a["minutos"],a["motivo"]))
        target=data.get("aplicar_para","todos"); ids=data.get("colaboradores",[])
        if target=="todos": cur.execute("SELECT id FROM colaboradores WHERE LOWER(COALESCE(status,'ativo'))='ativo'")
        elif target=="setor": cur.execute("SELECT id FROM colaboradores WHERE LOWER(COALESCE(status,'ativo'))='ativo' AND COALESCE(setor,'Sem setor')=%s",(data.get("setor"),))
        elif target=="selecionados" and ids: cur.execute("SELECT id FROM colaboradores WHERE id=ANY(%s)",(ids,))
        else: raise ValueError("Selecione colaboradores para aplicar a carga.")
        targets=[r[0] for r in cur.fetchall()]
        for cid in targets: cur.execute("INSERT INTO banco_horas_colaborador(periodo_id,colaborador_id,horas_previstas_minutos,atualizado_por) VALUES(%s,%s,%s,%s) ON CONFLICT(periodo_id,colaborador_id) DO UPDATE SET horas_previstas_minutos=EXCLUDED.horas_previstas_minutos,atualizado_por=EXCLUDED.atualizado_por,atualizado_em=NOW()",(pid,cid,final,user_id()))
        audit(cur,pid,None,"configurar_carga",old,final,data.get("observacao")); conn.commit()
        return jsonify({"success":True,"aplicados":len(targets),"prevista_minutos":final,"prevista":format_minutes(final)})
    except (ValueError,TypeError) as e: conn.rollback(); return jsonify({"error":str(e)}),400
    finally: cur.close(); conn.close()

@banco_horas_bp.post("/copiar-anterior")
def copy_previous():
    data=request.get_json(silent=True) or {}
    try: start=competence(data.get("competencia"))
    except ValueError as e: return jsonify({"error":str(e)}),400
    previous=(start.replace(day=1)-timedelta(days=1)).replace(day=1); conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("SELECT metodo_carga,carga_diaria_minutos,dias_considerados,carga_base_minutos,carga_prevista_padrao_minutos,observacao,id FROM banco_horas_periodos WHERE competencia=%s",(previous,)); source=cur.fetchone()
        if not source: return jsonify({"error":"O mês anterior não possui configuração."}),404
        pid=ensure_period(cur,start)
        cur.execute("UPDATE banco_horas_periodos SET metodo_carga=%s,carga_diaria_minutos=%s,dias_considerados=%s,carga_base_minutos=%s,carga_prevista_padrao_minutos=%s,observacao=%s,atualizado_em=NOW() WHERE id=%s",(*source[:6],pid))
        cur.execute("DELETE FROM banco_horas_ajustes WHERE periodo_id=%s",(pid,))
        cur.execute("INSERT INTO banco_horas_ajustes(periodo_id,minutos,motivo) SELECT %s,minutos,motivo FROM banco_horas_ajustes WHERE periodo_id=%s",(pid,source[6]))
        cur.execute("""INSERT INTO banco_horas_colaborador(periodo_id,colaborador_id,horas_previstas_minutos,atualizado_por)
                       SELECT %s,colaborador_id,horas_previstas_minutos,%s FROM banco_horas_colaborador
                       WHERE periodo_id=%s AND horas_previstas_minutos IS NOT NULL
                       ON CONFLICT(periodo_id,colaborador_id) DO UPDATE SET horas_previstas_minutos=EXCLUDED.horas_previstas_minutos,
                       atualizado_por=EXCLUDED.atualizado_por,atualizado_em=NOW()""",(pid,user_id(),source[6]))
        audit(cur,pid,None,"copiar_mes_anterior",previous,start); conn.commit(); return jsonify({"success":True})
    finally: cur.close(); conn.close()

@banco_horas_bp.post("/lancamento-mensal")
def monthly_entries():
    data=request.get_json(silent=True) or {}
    try: start=competence(data.get("competencia"))
    except ValueError as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        pid=ensure_period(cur,start)
        for item in data.get("lancamentos",[]):
            cid=int(item["colaborador_id"]); worked=parse_minutes(item.get("trabalhadas")); expected=parse_minutes(item["previstas"]) if item.get("previstas") not in (None,"") else None; present=max(0,int(item.get("presentes",0))); absent=max(0,int(item.get("ausentes",0)))
            cur.execute("SELECT horas_previstas_minutos,horas_trabalhadas_manual_minutos,dias_presentes_manual,dias_ausentes_manual FROM banco_horas_colaborador WHERE periodo_id=%s AND colaborador_id=%s",(pid,cid)); old=cur.fetchone()
            cur.execute("""INSERT INTO banco_horas_colaborador(periodo_id,colaborador_id,horas_previstas_minutos,origem_horas,horas_trabalhadas_manual_minutos,origem_frequencia,dias_presentes_manual,dias_ausentes_manual,atualizado_por) VALUES(%s,%s,%s,'manual',%s,'manual',%s,%s,%s) ON CONFLICT(periodo_id,colaborador_id) DO UPDATE SET horas_previstas_minutos=COALESCE(EXCLUDED.horas_previstas_minutos,banco_horas_colaborador.horas_previstas_minutos),origem_horas='manual',horas_trabalhadas_manual_minutos=EXCLUDED.horas_trabalhadas_manual_minutos,origem_frequencia='manual',dias_presentes_manual=EXCLUDED.dias_presentes_manual,dias_ausentes_manual=EXCLUDED.dias_ausentes_manual,atualizado_por=EXCLUDED.atualizado_por,atualizado_em=NOW()""",(pid,cid,expected,worked,present,absent,user_id())); audit(cur,pid,cid,"lancamento_mensal",old,(expected,worked,present,absent))
        conn.commit(); return jsonify({"success":True,"salvos":len(data.get("lancamentos",[]))})
    except (ValueError,TypeError,KeyError) as e: conn.rollback(); return jsonify({"error":str(e)}),400
    finally: cur.close(); conn.close()

def set_exclusion(excluded):
    data=request.get_json(silent=True) or {}
    try: start=competence(data.get("competencia")); ids=[int(x) for x in data.get("colaboradores",[])]; assert ids
    except (ValueError,TypeError,AssertionError): return jsonify({"error":"Período e colaboradores são obrigatórios."}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        pid=ensure_period(cur,start)
        for cid in ids:
            cur.execute("INSERT INTO banco_horas_colaborador(periodo_id,colaborador_id,excluido_periodo,motivo_exclusao,excluido_por,excluido_em,atualizado_por) VALUES(%s,%s,%s,%s,%s,CASE WHEN %s THEN NOW() END,%s) ON CONFLICT(periodo_id,colaborador_id) DO UPDATE SET excluido_periodo=EXCLUDED.excluido_periodo,motivo_exclusao=EXCLUDED.motivo_exclusao,excluido_por=EXCLUDED.excluido_por,excluido_em=EXCLUDED.excluido_em,atualizado_por=EXCLUDED.atualizado_por,atualizado_em=NOW()",(pid,cid,excluded,data.get("motivo") if excluded else None,user_id(),excluded,user_id())); audit(cur,pid,cid,"anular_periodo" if excluded else "restaurar_periodo",not excluded,excluded,data.get("motivo"))
        conn.commit(); return jsonify({"success":True,"alterados":len(ids)})
    finally: cur.close(); conn.close()

@banco_horas_bp.post("/anular")
def exclude(): return set_exclusion(True)
@banco_horas_bp.post("/restaurar")
def restore(): return set_exclusion(False)

@banco_horas_bp.get("/detalhe/<int:collaborator_id>")
def detail(collaborator_id):
    try: start=competence(request.args.get("competencia"))
    except ValueError as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        rows=get_summary(cur,start,{"colaborador_id":collaborator_id,"status":"todos"})
        if not rows: return jsonify({"error":"Colaborador não encontrado."}),404
        cur.execute("SELECT id,data_registro,hora_entrada,hora_saida,intervalo_minutos,situacao_dia,observacao FROM ponto WHERE colaborador_id=%s AND data_registro BETWEEN %s AND %s ORDER BY data_registro,hora_entrada",(collaborator_id,start,period_end(start)))
        records=[]
        for r in cur.fetchall():
            worked=calculate_record_minutes(r[2],r[3],r[4])
            records.append({"id":r[0],"data":r[1].isoformat(),"entrada":r[2].strftime('%H:%M') if r[2] else None,"saida":r[3].strftime('%H:%M') if r[3] else None,"intervalo_minutos":r[4],"trabalhadas":format_minutes(worked),"situacao":r[5],"observacao":r[6]})
        return jsonify({"colaborador":rows[0],"registros":records})
    finally: cur.close(); conn.close()

@banco_horas_bp.post("/registro")
def add_record():
    data=request.get_json(silent=True) or {}
    try:
        day=datetime.strptime(data["data"],"%Y-%m-%d").date(); start_time=datetime.combine(day,datetime.strptime(data["entrada"],"%H:%M").time()) if data.get("entrada") else None; end_time=datetime.combine(day,datetime.strptime(data["saida"],"%H:%M").time()) if data.get("saida") else None
        if start_time and end_time and end_time<start_time: end_time+=timedelta(days=1)
        interval=parse_minutes(data.get("intervalo","0")); situation=data.get("situacao","presente")
        if situation not in {'presente','ausente','ausencia_justificada','folga','ferias','afastamento'}: raise ValueError("Situação inválida.")
    except (ValueError,KeyError) as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("INSERT INTO ponto(colaborador_id,data_registro,hora_entrada,hora_saida,intervalo_minutos,situacao_dia,tipo_registro,observacao) VALUES(%s,%s,%s,%s,%s,%s,'manual',%s) RETURNING id",(data.get("colaborador_id"),day,start_time,end_time,interval,situation,data.get("observacao"))); rid=cur.fetchone()[0]; audit(cur,None,data.get("colaborador_id"),"adicionar_registro",None,rid,data.get("observacao")); conn.commit(); return jsonify({"success":True,"id":rid})
    finally: cur.close(); conn.close()

@banco_horas_bp.get("/exportar")
def export_excel():
    try: start=competence(request.args.get("competencia"))
    except ValueError as e: return jsonify({"error":str(e)}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        rows=get_summary(cur,start,request.args); include=request.args.get("incluir_anulados")=="1"
        if not include: rows=[r for r in rows if not r["excluido_periodo"]]
        wb=Workbook(); ws=wb.active; ws.title="Resumo"; ws.append(["Banco de Horas",start.strftime("%m/%Y"),"Gerado em",datetime.now().strftime("%d/%m/%Y %H:%M")]); ws.append([])
        headers=["Funcionário","Cargo","Setor","Horas Previstas","Horas Trabalhadas","Saldo","Horas Extras","Horas Negativas","Dias Presentes","Dias Ausentes","Origem das Horas","Situação","Status no período"] if include else ["Funcionário","Cargo","Setor","Horas Previstas","Horas Trabalhadas","Saldo","Horas Extras","Horas Negativas","Dias Presentes","Dias Ausentes","Origem das Horas","Situação"]
        ws.append(headers)
        for r in rows:
            values=[r["nome"],r["cargo"],r["setor"],r["previstas_minutos"]/1440,r["trabalhadas_minutos"]/1440,r["saldo"],r["extras_minutos"]/1440,r["negativas_minutos"]/1440,r["dias_presentes"],r["dias_ausentes"],r["origem_horas"].title(),r["situacao"]]
            if include: values.append("Anulado" if r["excluido_periodo"] else "Válido")
            ws.append(values)
        valid=[r for r in rows if not r["excluido_periodo"]]; totals=["TOTAL","","",sum(r["previstas_minutos"] for r in valid)/1440,sum(r["trabalhadas_minutos"] for r in valid)/1440,format_minutes(sum(r["saldo_minutos"] for r in valid),True),sum(r["extras_minutos"] for r in valid)/1440,sum(r["negativas_minutos"] for r in valid)/1440,sum(r["dias_presentes"] for r in valid),sum(r["dias_ausentes"] for r in valid),"",""]
        if include: totals.append("")
        ws.append(totals); header_row=3; ws.freeze_panes="A4"; ws.auto_filter.ref=f"A{header_row}:{get_column_letter(len(headers))}{max(header_row,len(rows)+3)}"
        for cell in ws[header_row]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="167052")
        for row in range(4,ws.max_row+1):
            for col in (4,5,7,8): ws.cell(row,col).number_format="[h]:mm"
        for i,width in enumerate((28,22,20,17,19,14,15,17,16,15,18,14,16),1): ws.column_dimensions[get_column_letter(i)].width=width
        detail_ws=wb.create_sheet("Detalhamento"); detail_ws.append(["Funcionário","Data","Entrada","Saída","Intervalo","Horas Trabalhadas","Presença","Observação"])
        ids=[r["id"] for r in rows]
        if ids:
            cur.execute("SELECT c.nome,p.data_registro,p.hora_entrada,p.hora_saida,p.intervalo_minutos,p.situacao_dia,p.observacao FROM ponto p JOIN colaboradores c ON c.id=p.colaborador_id WHERE p.colaborador_id=ANY(%s) AND p.data_registro BETWEEN %s AND %s ORDER BY c.nome,p.data_registro",(ids,start,period_end(start)))
            for r in cur.fetchall():
                worked=calculate_record_minutes(r[2],r[3],r[4])
                detail_ws.append([r[0],r[1],r[2].strftime('%H:%M') if r[2] else '',r[3].strftime('%H:%M') if r[3] else '',format_minutes(r[4]),format_minutes(worked),r[5],r[6]])
        detail_ws.freeze_panes="A2"; detail_ws.auto_filter.ref=detail_ws.dimensions
        config_ws=wb.create_sheet("Configuração"); config_ws.append(["Funcionário","Método da carga","Carga prevista final","Origem das horas","Status no período"])
        for r in rows: config_ws.append([r["nome"],"Individual/período",format_minutes(r["previstas_minutos"]),r["origem_horas"],"Anulado" if r["excluido_periodo"] else "Válido"])
        stream=io.BytesIO(); wb.save(stream); stream.seek(0); name=secure_filename(f"banco_horas_{start:%Y-%m}.xlsx")
        return send_file(stream,as_attachment=True,download_name=name,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally: cur.close(); conn.close()

@banco_horas_bp.get("/status")
def status(): return jsonify({"status":"ok","modelo":"consolidacao_mensal"})
